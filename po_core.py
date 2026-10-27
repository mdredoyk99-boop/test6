"""
po_core.py
==========
Core logic for the PO extractor: PDF parsing + Excel export.
Deliberately free of any UI code so it can be tested / reused from a script.

PDF layout this parser targets (WIP Sourcing Asia "Purchase Order"):

    Customer      Style name        Style No.   Color     XS S M L XL XXL  Total  Price  ETD  Ship
    Order No.
    <Style name> - <style code>                                             <- style header line
    26TRHK02215   I026392.00F.XX.03 Black / Gold  28 83 120 89 46 11  377   5.39  15 Oct 26 SEA
    ...
    Sub total                                                        14,574 68,363.21
    TOTAL                                                            96,735 454,801.85

Why coordinates instead of plain text?
    Blank size cells (e.g. no XS) simply vanish from extracted text, so
    "5 17 19 13 7" cannot be mapped to sizes by order alone. We therefore read
    word positions and assign every number to the nearest size-header column.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import BinaryIO, Optional

import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #
SIZES = ["XS", "S", "M", "L", "XL", "XXL"]

# Output columns A..N exactly as in "Actual PO.xlsx" (two "Color Code" / "Color"
# pairs exist in the template: B..E are data, O..P are lookup columns).
DATA_HEADERS = [
    "File", "Sty", "PO", "Color Code", "Color",
    "XS", "S", "M", "L", "XL", "XXL", "Pc", "Price", "ETD",
]
# Columns O..Y in the template are XLOOKUPs into an external workbook
# ('[1]SS''27'). They cannot be derived from the PDF, so we only reproduce the
# header cells (left empty) so the layout is identical.
LOOKUP_HEADERS = [
    "Color Code", "Color", "XS", "S", "M", "L", "XL", "XXL", "Pc", "Price", "ETD",
]

# The template joins style name and style code with 3 spaces:
#   "L/S Chase T-Shirt Loose   - 36N7.Q16"
STYLE_JOINER = "   "

# Regexes ------------------------------------------------------------------ #
# Customer order no.: 26TRHK02215 / 26TR004689 / 26TRUS04688 ...
RE_ORDER = re.compile(r"^\d{2}[A-Z]{2,}[A-Z]*\d+$")
# Style no.: I026392.00F.XX.03  (letter/digits separated by dots, >= 3 parts)
RE_STYLE_NO = re.compile(r"^[A-Z0-9]+(?:\.[A-Z0-9]+){2,}$")
# "L/S Chase T-Shirt Loose - 36N7.Q16" -> name / code
RE_STYLE_HEADER = re.compile(r"^(?P<name>.+?)\s+-\s+(?P<code>\S+)$")
RE_INT = re.compile(r"^\d{1,3}(?:,\d{3})*$|^\d+$")
RE_DECIMAL = re.compile(r"^\d{1,3}(?:,\d{3})*\.\d+$|^\d+\.\d+$")
RE_ETD = re.compile(r"(\d{1,2})\s+([A-Za-z]{3})\s+(\d{2,4})")


# --------------------------------------------------------------------------- #
# Data containers
# --------------------------------------------------------------------------- #
@dataclass
class ParseResult:
    """Result of parsing one PDF."""
    file_label: str
    rows: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    printed_total_pcs: Optional[int] = None      # "TOTAL" line on the PO
    printed_total_amount: Optional[float] = None


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def file_label_from_name(filename: str) -> str:
    """
    Convert an uploaded filename into the label used in column A.

        WBDCSBC271740_20_.pdf   -> WBDCSBC271740[20]
        WBDCSBC271740 (20).pdf  -> WBDCSBC271740[20]
        WBDCSBC271840.pdf       -> WBDCSBC271840
    """
    stem = re.sub(r"\.pdf$", "", filename.strip(), flags=re.I)
    m = re.match(r"^(.*?)_(\d+)_$", stem) or re.match(r"^(.*?)\s*\((\d+)\)$", stem)
    return f"{m.group(1)}[{m.group(2)}]" if m else stem


def _to_int(text: str) -> int:
    return int(text.replace(",", ""))


def _to_float(text: str) -> float:
    return float(text.replace(",", ""))


def _group_lines(words: list[dict], tol: float = 3.5) -> list[list[dict]]:
    """Group pdfplumber words into visual lines (by vertical position)."""
    lines: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(l, key=lambda w: w["x0"]) for l in lines]


def _center(w: dict) -> float:
    return (w["x0"] + w["x1"]) / 2


def _line_text(line: list[dict]) -> str:
    return " ".join(w["text"] for w in line)


def _find_size_header(lines: list[list[dict]]) -> Optional[tuple[dict, float]]:
    """
    Locate the header line holding XS..XXL and return
    ({size: x_center}, column_pitch). Returns None if the page has no header.
    """
    for line in lines:
        texts = [w["text"] for w in line]
        if "XS" in texts and "XXL" in texts:
            centers = {}
            for w in line:
                if w["text"] in SIZES and w["text"] not in centers:
                    centers[w["text"]] = _center(w)
            if len(centers) == len(SIZES):
                xs = [centers[s] for s in SIZES]
                pitch = (xs[-1] - xs[0]) / (len(xs) - 1)
                return centers, pitch
    return None


def _parse_etd(text: str) -> Optional[datetime]:
    m = RE_ETD.search(text)
    if not m:
        return None
    day, mon, year = m.groups()
    fmt = "%d %b %y" if len(year) == 2 else "%d %b %Y"
    try:
        return datetime.strptime(f"{day} {mon.title()} {year}", fmt)
    except ValueError:
        return None


# --------------------------------------------------------------------------- #
# Main PDF parser
# --------------------------------------------------------------------------- #
def parse_po_pdf(source: BinaryIO | str, filename: str) -> ParseResult:
    """
    Parse one PO PDF into item rows.

    `source` may be a path or a binary file-like object (Streamlit upload).
    Only genuine item rows are returned: header blocks, style headings,
    Sub total / TOTAL lines, remarks and signature areas are all ignored.
    """
    result = ParseResult(file_label=file_label_from_name(filename))

    current_style: Optional[str] = None     # carried across pages if a style
    last_header: Optional[tuple] = None     # block spans a page break

    with pdfplumber.open(source) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            lines = _group_lines(words)

            header = _find_size_header(lines) or last_header
            if header is None:
                # Page without a table header and none seen before (e.g. a
                # cover / terms page) -> nothing to extract.
                continue
            last_header = header
            centers, pitch = header
            size_names = SIZES
            first_size_x = centers["XS"]
            last_size_x = centers["XXL"]
            size_zone_end = last_size_x + pitch / 2     # right edge of XXL col
            color_zone_end = first_size_x - pitch / 2   # left edge of XS col

            in_table = False
            for line in lines:
                text = _line_text(line)
                first = line[0]["text"]

                # ---- table boundaries -------------------------------------
                if "XS" in [w["text"] for w in line] and "XXL" in [w["text"] for w in line]:
                    in_table = True
                    continue
                if not in_table:
                    continue

                # ---- Sub total / TOTAL lines (used only for reconciliation)
                if first.lower() == "sub" or first.upper() == "TOTAL":
                    if first.upper() == "TOTAL":
                        nums = [w["text"] for w in line[1:]]
                        if len(nums) >= 2 and RE_INT.match(nums[0]) and RE_DECIMAL.match(nums[1]):
                            result.printed_total_pcs = _to_int(nums[0])
                            result.printed_total_amount = _to_float(nums[1])
                    continue
                if first.startswith("Remarks"):
                    break  # nothing after this is table data

                # ---- item row --------------------------------------------
                if RE_ORDER.match(first) and len(line) > 1 and RE_STYLE_NO.match(line[1]["text"]):
                    row = _parse_item_line(line, current_style, size_names, centers,
                                           color_zone_end, size_zone_end, result, page_no)
                    if row:
                        row["File"] = result.file_label
                        result.rows.append(row)
                    continue

                # ---- style heading ("L/S Chase T-Shirt Loose - 36N7.Q16") --
                m = RE_STYLE_HEADER.match(text)
                if m and not any(RE_INT.match(w["text"]) for w in line):
                    current_style = f"{m.group('name')}{STYLE_JOINER}- {m.group('code')}"

    if not result.rows:
        result.warnings.append("No item rows found - is this the expected PO layout?")
    _reconcile(result)
    return result


def _parse_item_line(line, style, size_names, centers, color_zone_end,
                     size_zone_end, result: ParseResult, page_no: int) -> Optional[dict]:
    """Turn one visual line of the table into a row dict."""
    order_no = line[0]["text"]
    style_no = line[1]["text"]
    rest = line[2:]

    # Colour = everything left of the XS column (digits in a colour name are safe).
    color = " ".join(w["text"] for w in rest if _center(w) < color_zone_end)

    # Sizes: numbers inside the size zone -> nearest header column.
    sizes = {s: 0 for s in size_names}
    seen_blank_collision = False
    for w in rest:
        if color_zone_end <= _center(w) < size_zone_end and RE_INT.match(w["text"]):
            nearest = min(size_names, key=lambda s: abs(centers[s] - _center(w)))
            if sizes[nearest]:
                seen_blank_collision = True
            sizes[nearest] = _to_int(w["text"])
    if seen_blank_collision:
        result.warnings.append(f"{order_no}: two numbers mapped to the same size column (page {page_no}).")

    # Tail: Total, Unit price, ETD (3 words), Ship mode.
    tail = [w["text"] for w in rest if _center(w) >= size_zone_end]
    total = price = etd = None
    ship = ""
    if len(tail) >= 2 and RE_INT.match(tail[0]) and RE_DECIMAL.match(tail[1]):
        total, price = _to_int(tail[0]), _to_float(tail[1])
        etd = _parse_etd(" ".join(tail[2:5]))
        ship = " ".join(tail[5:])
    else:
        result.warnings.append(f"{order_no}: could not read Total/Price columns (page {page_no}): {tail}")
        return None

    if sum(sizes.values()) != total:
        result.warnings.append(
            f"{order_no}: size breakdown sums to {sum(sizes.values())} but Total is {total}.")
    if etd is None:
        result.warnings.append(f"{order_no}: ETD not recognised in {tail[2:5]}.")

    return {
        "Sty": style or "",
        "PO": order_no,
        "Color Code": style_no,
        "Color": color,
        **sizes,
        "Pc": total,
        "Price": price,
        "ETD": etd,
        "_ship": ship,     # not exported (template has no Ship Mode column)
    }


def _reconcile(result: ParseResult) -> None:
    """Compare parsed rows with the TOTAL line printed on the PO."""
    if result.printed_total_pcs is None:
        result.warnings.append("Could not find the TOTAL line to reconcile against.")
        return
    pcs = sum(r["Pc"] for r in result.rows)
    amount = round(sum(r["Pc"] * r["Price"] for r in result.rows), 2)
    if pcs != result.printed_total_pcs:
        result.warnings.append(
            f"Piece count mismatch: parsed {pcs:,} vs printed TOTAL {result.printed_total_pcs:,}.")
    if abs(amount - (result.printed_total_amount or 0)) > 1.0:
        result.warnings.append(
            f"Amount mismatch: parsed {amount:,.2f} vs printed TOTAL {result.printed_total_amount:,.2f}.")


# --------------------------------------------------------------------------- #
# DataFrame + Excel export
# --------------------------------------------------------------------------- #
def build_dataframe(results: list[ParseResult], blank_sizes_as_zero: bool = True) -> pd.DataFrame:
    """Combine results into a DataFrame whose columns mirror 'Actual PO.xlsx' A..N."""
    rows = [{h: r.get(h) for h in DATA_HEADERS} for res in results for r in res.rows]
    df = pd.DataFrame(rows, columns=DATA_HEADERS)
    if not blank_sizes_as_zero:
        # Reference sheet leaves some blank PDF cells empty rather than 0.
        for s in SIZES:
            df[s] = df[s].astype("object").where(df[s] != 0, None)
    return df


# Column widths copied from the template.
_COL_WIDTHS = {"A": 20.18, "B": 38.63, "C": 12.54, "D": 17.54, "E": 34.63,
               "F": 5.91, "G": 5.36, "J": 5.73, "K": 6.82, "L": 6.36,
               "M": 7.45, "N": 9.27, "O": 17.73, "P": 35.0, "Y": 10.63}
# Header fonts copied from the template: (name, size, bold)
_HEADER_FONT = {"A": ("Aptos Narrow", 11, False), "B": ("Arial", 10, True),
                "C": ("Arial", 10, False), "D": ("Arial", 10, True),
                "E": ("Arial", 10, True), "O": ("Arial", 10, True),
                "P": ("Arial", 10, True)}


def to_excel_bytes(df: pd.DataFrame, include_lookup_headers: bool = True) -> bytes:
    """
    Write the DataFrame to an .xlsx that reproduces the template's layout:
    header row, Arial 10 left-aligned data, bold style column, 'd-mmm-yy' ETD,
    '#,##0' piece counts, autofilter and column widths.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = DATA_HEADERS + (LOOKUP_HEADERS if include_lookup_headers else [])
    left = Alignment(horizontal="left")

    # Header row
    for i, h in enumerate(headers, start=1):
        col = get_column_letter(i)
        name, size, bold = _HEADER_FONT.get(col, ("Arial", 9, False))
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=name, size=size, bold=bold)
        c.alignment = left

    # Data rows (columns A..N)
    for r_idx, rec in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(rec, start=1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            elif not isinstance(value, str) and pd.isna(value):
                value = None                      # NaN/None -> truly empty cell
            c = ws.cell(row=r_idx, column=c_idx, value=value)
            c.font = Font(name="Arial", size=10, bold=(c_idx == 2))   # col B bold
            c.alignment = left
            if DATA_HEADERS[c_idx - 1] == "ETD":
                c.number_format = "d-mmm-yy"
            elif DATA_HEADERS[c_idx - 1] == "Pc":
                c.number_format = "#,##0"

    for col, width in _COL_WIDTHS.items():
        if include_lookup_headers or col <= "N":
            ws.column_dimensions[col].width = width
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{max(ws.max_row, 2)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
